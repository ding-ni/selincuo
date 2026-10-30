from __future__ import annotations

import calendar
import hashlib
import json
import math
import re
import shutil
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from .common import ROOT, canonical, now, sha_file, write_json
from .database import Database

STATIONS = [
    ('BLQ', '巴勒切', 'rain', 91.87776607, 32.69167210),
    ('115DB', '115道班', 'rain', 91.83398992, 32.51681861),
    ('DUJA', '杜加', 'river', 89.15997904, 32.03057137),
    ('SLC', '色林错主站', 'lake', 89.35702225, 31.84329578),
    ('SLC-E', '色林错东岸', 'lake', None, None),
    ('SLC-S', '色林错南岸', 'lake', None, None),
    ('SLC-W', '色林错西岸', 'lake', None, None),
    ('SLC-N', '色林错北岸', 'lake', None, None),
]


def shore_station(sheet):
    """Map the monitoring system's shore sheet names to their own stations."""
    title = re.sub(r'[\s_\-（）()]+', '', str(sheet))
    match = re.fullmatch(r'(?:色林错|色|\d{2,4})?([东南西北])(?:岸|站)?(?:水位)?(?:数据|资料|表)?', title)
    if not match:
        return None
    return 'SLC-' + {'东': 'E', '南': 'S', '西': 'W', '北': 'N'}[match.group(1)]


def station_for(path, sheet=''):
    parts = str(path).replace('\\', '/').lower()
    if '115daoban/' in parts:
        return '115DB'
    if 'blq/' in parts:
        return 'BLQ'
    if 'dj/' in parts:
        return 'DUJA'
    if '东南西北' in parts:
        station = shore_station(sheet)
        if station:
            return station
        raise ValueError(f'四岸水位工作表缺少明确方位：{sheet}')
    return 'SLC'


def variable_for(path):
    text = str(path)
    if '雨量' in text:
        return 'precipitation'
    if '流量' in text:
        return 'flow'
    if '水位' in text:
        return 'level'
    return None


def number(value):
    if isinstance(value, bool):
        return None
    try:
        result = float(str(value).strip())
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None


def expand_level(raw, previous):
    text = str(raw).strip()
    qualified = text.startswith('(') or text.startswith('（')
    value = number(text.strip('()（） '))
    if value is None:
        return None, previous, 'ice_or_stopped' if text in ('L', 'T', '连底冻', '〃') else 'missing'
    if value.is_integer() and abs(value) < 100 and previous is not None:
        value = math.floor(previous) + value / 100
    if value < 4000 or value > 6500:
        return None, previous, 'unresolved_level_abbreviation'
    return value, value, 'qualified' if qualified else 'valid'


def scan(root):
    root = Path(root)
    return [{'path': str(p.relative_to(root)).replace('\\', '/'), 'bytes': p.stat().st_size,
             'sha256': sha_file(p), 'kind': variable_for(p), 'status': 'ready' if variable_for(p) else 'station-metadata'}
            for p in sorted(root.rglob('*')) if p.suffix.lower() in ('.xls', '.xlsx') and not p.name.startswith('~$')]


def xlsx_records(path, station_override=None, variable_override=None):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    variable = variable_override or variable_for(path)
    for sheet in wb:
        if station_override == 'SHORES':
            station = shore_station(sheet.title)
            if station is None:
                continue
        else:
            station = station_for(path, sheet.title) if station_override is None else station_override
        previous_level = None
        for index, values in enumerate(sheet.iter_rows(values_only=True), 1):
            if not values or not isinstance(values[0], datetime):
                continue
            time = values[0].replace(tzinfo=None)
            raw = values[-1]
            value = number(raw)
            quality = 'valid' if value is not None else ('ice_or_stopped' if str(raw).strip() in ('L', 'T') else 'missing')
            if variable == 'level':
                value, previous_level, quality = expand_level(raw, previous_level)
            yield {'station': station, 'variable': variable, 'time': time.isoformat(timespec='seconds'),
                   'start': None, 'end': time.isoformat(timespec='seconds'), 'step': 'irregular',
                   'value': value, 'raw': str(raw), 'quality': quality,
                   'sheet': sheet.title, 'cell': f'R{index}C{len(values)}'}
    wb.close()


def _xls_report_year(wb):
    """Read only an explicitly labelled report year, not its creation date."""
    years = set()
    for sheet in wb.sheets():
        for r in range(min(8, sheet.nrows)):
            row = [str(sheet.cell_value(r, c)).strip() for c in range(min(16, sheet.ncols))]
            for index, cell in enumerate(row):
                if '年份' not in cell and '年度' not in cell:
                    continue
                for text in row[index:min(len(row), index + 4)]:
                    match = re.search(r'(?<!\d)(20\d{2})(?!\d)', text)
                    if match:
                        years.add(int(match.group(1)))
                        break
    if len(years) > 1:
        raise ValueError('工作表中的年度不一致，请核对原表')
    return next(iter(years), None)


def xls_records(path, station_override=None, variable_override=None, source_name=None):
    import xlrd
    wb = xlrd.open_workbook(str(path), on_demand=True)
    variable = variable_override or variable_for(path)
    # Incoming files have random prefixes that can themselves contain 20xx.
    # Only the user's original file name (or a labelled report year) is data.
    name = Path(path).name if source_name is None else Path(source_name).name
    match = re.search(r'(?<!\d)(20\d{2})(?!\d)', name)
    name_year = int(match.group(1)) if match else None
    report_year = _xls_report_year(wb)
    if name_year and report_year and name_year != report_year:
        wb.release_resources()
        raise ValueError('文件名年度与工作表年度不一致，请核对原表')
    year = name_year or report_year
    if not year or not variable:
        wb.release_resources()
        return
    for sheet in wb.sheets():
        if station_override == 'SHORES':
            station = shore_station(sheet.name)
            if station is None:
                continue
        else:
            station = station_for(path, sheet.name) if station_override is None else station_override
        monthly_headers = None
        for r in range(min(8, sheet.nrows)):
            row = [str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)]
            if '一月' in row and '十二月' in row:
                monthly_headers = (r, row.index('一月'))
                break
        if monthly_headers is None:
            continue
        header, first_month = monthly_headers
        for month in range(1, 13):
            column = first_month + month - 1
            previous = None
            inherited_flag = None
            for r in range(header + 1, min(sheet.nrows, header + 35)):
                day = number(sheet.cell_value(r, 0))
                if day is None or int(day) != day or not 1 <= day <= calendar.monthrange(year, month)[1]:
                    continue
                raw = str(sheet.cell_value(r, column)).strip()
                time = datetime(year, month, int(day))
                quality = 'valid'
                if variable == 'level':
                    if raw == '〃' and inherited_flag:
                        value, quality = None, inherited_flag
                    else:
                        value, previous, quality = expand_level(raw, previous)
                    inherited_flag = quality if quality != 'valid' else None
                else:
                    value = number(raw.strip('()（） '))
                    quality = 'qualified' if '(' in raw or '（' in raw else 'valid'
                    if value is None:
                        quality = 'missing'
                    elif variable == 'flow' and value == 0:
                        quality = 'zero_needs_review'
                yield {'station': station, 'variable': variable, 'time': time.isoformat(timespec='seconds'),
                       'start': time.isoformat(timespec='seconds'), 'end': (time + timedelta(days=1)).isoformat(timespec='seconds'),
                       'step': 'day', 'value': value, 'raw': raw, 'quality': quality,
                       'sheet': sheet.name, 'cell': f'R{r + 1}C{column + 1}'}
    wb.release_resources()


def import_observations(database=None, root=None, progress=None):
    db = database or Database()
    root = Path(root or ROOT / 'data/raw/observations/slc')
    files = scan(root)
    with db.connect() as conn:
        for station, name, kind, lon, lat in STATIONS:
            conn.execute('INSERT OR IGNORE INTO stations VALUES(?,?,?,?,?,?)',
                         (station, name, kind, lon, lat, canonical({'datumVerified': False, 'projectId': 'selinco'})))
    from .stations import sync_locations, sync_datums
    sync_locations(db)
    sync_datums(db)
    report = {'createdAt': now(), 'files': [], 'rainfallPolicy': 'interval-increment; absent=zero ONLY within registered observation coverage',
              'dailyStatisticalBoundary': 'source-labelled day; precipitation day-boundary must be checked before correction'}
    for index, info in enumerate(files):
        if progress:
            progress((index / max(len(files), 1)) * 100, f"导入 {info['path']}")
        path = root / info['path']
        archived = db.workspace / 'observation-sources' / info['sha256'] / path.name
        archived.parent.mkdir(parents=True, exist_ok=True)
        if not archived.is_file():
            shutil.copy2(path, archived)
        info['archivedPath'] = archived.relative_to(db.workspace).as_posix()
        pieces = Path(info['path']).parts
        if 'imports' in pieces:
            logical_path = '/'.join([*pieces[:pieces.index('imports')], pieces[-1]])
        else:
            logical_path = info['path']
        source_id = 'SRC-' + hashlib.sha256((logical_path + info['sha256']).encode()).hexdigest()[:24]
        existing = db.rows('SELECT id,metadata FROM sources WHERE sha256=?', (info['sha256'],))
        if existing:
            source_id = existing[0]['id']
            with db.connect() as conn:
                conn.execute('UPDATE sources SET metadata=? WHERE id=?', (canonical({**info, 'parserVersion': 2}), source_id))
            report['files'].append({**info, 'sourceId': source_id, 'status': 'unchanged'})
            continue
        records = list(xlsx_records(path) if path.suffix.lower() == '.xlsx' else xls_records(path))
        groups = defaultdict(list)
        for record in records:
            groups[(record['station'], record['variable'], record['step'])].append(record)
        with db.connect() as conn:
            conn.execute('INSERT INTO sources VALUES(?,?,?,?,?,?)', (source_id, logical_path, info['sha256'], info['bytes'], now(), canonical({**info, 'parserVersion': 2})))
            for r in records:
                unit = {'precipitation': 'mm', 'flow': 'm3/s', 'level': 'm'}[r['variable']]
                conn.execute('''INSERT INTO observations(source_id,station_id,variable,time,period_start,period_end,step,unit,value,raw_text,quality,sheet,cell) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                             (source_id, r['station'], r['variable'], r['time'], r['start'], r['end'], r['step'], unit, r['value'], r['raw'], r['quality'], r['sheet'], r['cell']))
            for (station, variable, step), rows in groups.items():
                times = [r['time'] for r in rows]
                metadata = {'zeroForAbsent': variable == 'precipitation', 'coverageConfirmed': False,
                            'note': '稀疏雨量在确认的观测期内补零；2026年不从最后一场雨推断文件截止日。',
                            'flags': sorted({r['quality'] for r in rows})}
                conn.execute('INSERT INTO coverage VALUES(?,?,?,?,?,?,?,?,?)',
                             (source_id, station, variable, step, min(times), max(times), sum(r['value'] is not None for r in rows), len(rows), canonical(metadata)))
            # Invalidate observation views in the same transaction as the source.
            conn.execute('INSERT INTO audit(at,action,payload) VALUES(?,?,?)',
                         (now(), 'observations.source-import', canonical({'sourceId': source_id, 'records': len(records)})))
        report['files'].append({**info, 'sourceId': source_id, 'records': len(records), 'status': 'imported'})
    write_json(db.workspace / 'reports/observation-import.json', report)
    db.log('observations.import', {'files': len(files), 'imported': sum(x['status'] == 'imported' for x in report['files'])})
    return report


def daily_series(db, station, variable, start=None, end=None):
    """Prefer source daily means; aggregate subdaily levels by elapsed-time coverage."""
    import numpy as np
    import pandas as pd
    rows = db.observations(station, variable, start=start, end=end, limit=None)
    if not rows and variable != 'precipitation':
        return pd.Series(dtype=float)
    frame = pd.DataFrame(rows) if rows else pd.DataFrame(columns=['time', 'step', 'quality', 'value'])
    frame['date'] = pd.to_datetime(frame.time).dt.floor('D')
    daily = frame[(frame.step == 'day') & (frame.quality.isin(['valid', 'qualified']))].dropna(subset=['value'])
    result = daily.groupby('date').value.last()
    # Standard exchange hourly records mark the END of their interval. Midnight
    # belongs to the preceding day, and partial days cannot imply daily totals.
    hourly_rows = frame[(frame.step == 'hour') & frame.quality.eq('valid')].copy()
    if len(hourly_rows):
        hourly_rows.index = pd.to_datetime(hourly_rows.time)
        hourly_values = hourly_rows.value.sort_index().resample('D', closed='right', label='left')
        hourly_daily = (hourly_values.sum(min_count=24) if variable in ('precipitation', 'evaporation')
                        else hourly_values.mean().where(hourly_values.count() == 24))
        result = result.combine_first(hourly_daily)
    sub = frame[(~frame.step.isin(['day', 'hour'])) & frame.quality.eq('valid')].dropna(subset=['value']).copy()
    if variable == 'precipitation':
        # Each event source needs its own confirmed, complete observation day.
        # A positive event alone does not prove that the rest of the day was
        # observed; use one source per day and never add overlapping sources.
        coverages = db.rows('''SELECT c.* FROM coverage c JOIN sources s ON s.id=c.source_id
            WHERE c.station_id=? AND c.variable=? AND c.step='irregular'
            AND s.rowid IN (SELECT MAX(rowid) FROM sources GROUP BY path)
            ORDER BY s.imported_at DESC,s.rowid DESC''', (station, variable))
        event_daily = pd.Series(dtype=float)
        for c in coverages:
            metadata = json.loads(c['metadata'])
            if not metadata.get('coverageConfirmed') or not metadata.get('zeroForAbsent'):
                continue
            a, b = pd.Timestamp(c['start']), pd.Timestamp(c['end'])
            days = pd.date_range(a.ceil('D'), b.floor('D') - pd.Timedelta(days=1), freq='D')
            if start:
                days = days[days >= pd.Timestamp(start).floor('D')]
            if end:
                days = days[days <= pd.Timestamp(end).floor('D')]
            if not len(days):
                continue
            # Retain the confirmed coverage axis even when its final days are
            # stopped. A stopped day is missing, never a zero-rain day.
            operational = np.ones(len(days), dtype=bool)
            for stopped_start, stopped_end in metadata.get('stoppedIntervals', []):
                operational &= ((days + pd.Timedelta(days=1) <= pd.Timestamp(stopped_start)) |
                                (days >= pd.Timestamp(stopped_end)))
            source_rows = db.rows('''SELECT time,value,quality FROM observations
                WHERE source_id=? AND station_id=? AND variable=? AND step='irregular'
                ORDER BY time,id''', (c['source_id'], station, variable))
            # Native increments update individual event times without claiming
            # new coverage. Overlay them only inside this already-confirmed
            # coverage; missing days outside it must never become zero.
            increments = db.rows('''SELECT * FROM (SELECT o.time,o.value,o.quality,
                    ROW_NUMBER() OVER(PARTITION BY o.time ORDER BY s.imported_at DESC,o.id DESC) rn
                    FROM observations o JOIN sources s ON s.id=o.source_id
                    WHERE o.station_id=? AND o.variable=? AND o.step='irregular'
                    AND o.time>=? AND o.time<? AND s.path LIKE 'native-excel/%'
                    AND s.rowid>(SELECT rowid FROM sources WHERE id=?)) WHERE rn=1''',
                    (station, variable, a.isoformat(), b.isoformat(), c['source_id']))
            if increments:
                merged = {r['time']: r for r in source_rows}
                merged.update({r['time']: r for r in increments})
                source_rows = list(merged.values())
            values = pd.Series(0., index=days)
            if source_rows:
                source = pd.DataFrame(source_rows)
                source['date'] = pd.to_datetime(source.time).dt.floor('D')
                bad = source[(~source.quality.isin(['valid', 'qualified'])) |
                             source.value.isna() | (source.value < 0)]
                valid = source[source.quality.isin(['valid', 'qualified']) &
                               source.value.notna() & (source.value >= 0)]
                sums = valid.groupby('date').value.sum()
                common = values.index.intersection(sums.index)
                values.loc[common] = sums.loc[common]
                values.loc[values.index.intersection(bad.date)] = np.nan
            values.loc[days[~operational]] = np.nan
            event_daily = event_daily.combine_first(values)
        result = result.combine_first(event_daily)
    elif len(sub):
        sub.index = pd.to_datetime(sub.time)
        sub = sub.sort_index()
        hourly = sub.value.resample('1h').mean()
        counts = hourly.resample('D').count()
        agg = hourly.resample('D').mean().where(counts >= 12)
        result = result.combine_first(agg)
    result = result.sort_index()
    if start:
        result = result[result.index >= pd.Timestamp(start).floor('D')]
    if end:
        result = result[result.index <= pd.Timestamp(end).floor('D')]
    return result.astype(float)


def confirm_coverage(db, value):
    import pandas as pd
    a, b = pd.Timestamp(value['start']), pd.Timestamp(value['end'])
    if a >= b or a.tzinfo or b.tzinfo:
        raise ValueError('覆盖期起止无效，请使用北京时间')
    stopped = []
    for row in value.get('stops', '').splitlines():
        if not row.strip():
            continue
        x, y = [pd.Timestamp(s.strip()) for s in row.split(',')]
        if not a <= x < y <= b:
            raise ValueError('停测区间须位于覆盖期内')
        stopped.append([x.isoformat(), y.isoformat()])
    records = db.rows('SELECT * FROM coverage WHERE source_id=? AND station_id=? AND variable=?',
                      (value['sourceId'], value['station'], value['variable']))
    if not records:
        raise ValueError('来源、站点或变量未匹配')
    with db.connect() as conn:
        for r in records:
            metadata = {**json.loads(r['metadata']), 'coverageConfirmed': True, 'stoppedIntervals': stopped,
                        'confirmedAt': now(), 'note': str(value.get('note', '')), 'zeroForAbsent': value['variable'] == 'precipitation',
                        **({'eventTimeMeaning': value['eventTimeMeaning']} if value.get('eventTimeMeaning') else {})}
            conn.execute('UPDATE coverage SET start=?,end=?,metadata=? WHERE source_id=? AND station_id=? AND variable=? AND step=?',
                         (a.isoformat(), b.isoformat(), canonical(metadata), r['source_id'], r['station_id'], r['variable'], r['step']))
        conn.execute('INSERT INTO audit(at,action,payload) VALUES(?,?,?)',
                     (now(), 'coverage.confirm', canonical({**value, 'stoppedIntervals': stopped})))
    return {'status': 'confirmed', 'start': a.isoformat(), 'end': b.isoformat()}
