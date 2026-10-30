"""Incremental import of the station system's original Excel exports.

Every uploaded workbook is an immutable source. Earlier same-name uploads stay
active, so a later partial workbook cannot erase their non-overlapping dates.
"""
from collections import defaultdict
from pathlib import Path
import shutil

from .common import canonical, digest, now, sha_file, within
from .observations import shore_station, xls_records, xlsx_records, variable_for


STATIONS = {'115道班': '115DB', '巴勒切': 'BLQ', '杜加': 'DUJA', '色林错': 'SLC'}
UNITS = {'precipitation': 'mm', 'flow': 'm3/s', 'level': 'm'}


def _workbook_hints(path):
    """Inspect sheet names and report headers, never the generated upload name."""
    titles, cells = [], []
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book:
                titles.append(sheet.title)
                for row in sheet.iter_rows(min_row=1, max_row=8, max_col=16, values_only=True):
                    cells.extend(str(value) for value in row if isinstance(value, str))
        finally:
            book.close()
    else:
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            for sheet in book.sheets():
                titles.append(sheet.name)
                for r in range(min(8, sheet.nrows)):
                    cells.extend(str(sheet.cell_value(r, c)) for c in range(min(16, sheet.ncols)))
        finally:
            book.release_resources()
    text = ' '.join(titles + cells)
    return {'shoreSheets': [title for title in titles if shore_station(title)],
            'stations': {station for name, station in STATIONS.items() if name in text},
            'variables': {variable_for(text) for text in titles + cells if variable_for(text)}}


def _parse(db, item):
    if not isinstance(item, dict):
        raise ValueError('资料文件信息无效')
    path = within(db.workspace / 'incoming', item.get('filename', ''))
    original = item.get('originalName', '')
    if Path(original).name != original or not original or path.suffix.lower() not in ('.xls', '.xlsx'):
        raise ValueError('请选择原格式 Excel 文件')
    if not path.is_file():
        raise ValueError('上传文件不存在，请重新选择')
    hints = _workbook_hints(path)
    inferred = next((station for name, station in STATIONS.items() if name in original), None)
    if inferred is None and len(hints['stations']) == 1:
        inferred = next(iter(hints['stations']))
    station = item.get('station') or inferred
    shores = bool(hints['shoreSheets']) or '东南西北' in original
    if shores:
        station = 'SHORES'
    if station not in ('SLC', 'DUJA', 'BLQ', '115DB', 'SHORES'):
        raise ValueError(f'{original} 无法识别站点，请指定站点')
    variable = variable_for(original)
    if variable is None and len(hints['variables']) == 1:
        variable = next(iter(hints['variables']))
    if variable is None and shores:
        variable = 'level'
    if variable not in UNITS:
        raise ValueError(f'{original} 无法识别水位、流量或雨量，请使用原始导出文件名')
    if shores and variable != 'level':
        raise ValueError('四岸工作表仅支持湖泊水位')
    parser = xlsx_records if path.suffix.lower() == '.xlsx' else xls_records
    kwargs = {'station_override': station, 'variable_override': variable}
    if path.suffix.lower() == '.xls':
        kwargs['source_name'] = original
    rows = list(parser(path, **kwargs))
    if not rows:
        raise ValueError(f'{original} 未识别到观测记录，请检查工作表格式及年份')
    return path, original, station, variable, rows


def _active(conn, keys):
    result = {}
    for station, variable in {(k[0], k[1]) for k in keys}:
        rows = conn.execute('''SELECT * FROM (SELECT o.*,ROW_NUMBER() OVER(
            PARTITION BY station_id,variable,step,time ORDER BY s.imported_at DESC,o.id DESC) rn
            FROM observations o JOIN sources s ON s.id=o.source_id
            WHERE o.station_id=? AND o.variable=? AND s.rowid IN
            (SELECT MAX(rowid) FROM sources GROUP BY path)) WHERE rn=1''', (station, variable))
        for row in rows:
            r = dict(row)
            result[(r['station_id'], r['variable'], r['step'], r['time'])] = r
    return result


def _prepare(db, request, conn):
    files = request.get('files')
    if not isinstance(files, list) or not 1 <= len(files) <= 50:
        raise ValueError('请选择1至50个Excel文件')
    parsed = [_parse(db, item) for item in files]
    combined = {}
    for index, (_, name, _, _, rows) in enumerate(parsed):
        for row in rows:
            key = (row['station'], row['variable'], row['step'], row['time'])
            if key in combined:
                prior = combined[key][1]
                if all(prior[field] == row[field] for field in ('value', 'quality', 'start', 'end')):
                    continue
                raise ValueError(f'{name} 与本次其他表的同一时次数值不同，请分批核对')
            combined[key] = (index, row)
    existing = _active(conn, combined)
    counts = {'new': 0, 'changed': 0, 'same': 0, 'missing': 0, 'preserved': 0}
    changes, issues = [], []
    groups = defaultdict(list)
    for key, (index, row) in combined.items():
        prior = existing.get(key)
        unit = UNITS[row['variable']]
        same = prior is not None and all((prior[field] == value) for field, value in (
            ('value', row['value']), ('quality', row['quality']), ('unit', unit),
            ('period_start', row['start']), ('period_end', row['end'])))
        # Blank cells in a partial annual workbook are not deletion requests.
        # Explicit stop/ice markers remain real observations and may replace a value.
        preserve = prior is not None and prior['value'] is not None and row['quality'] == 'missing'
        status = 'preserved' if preserve else 'new' if prior is None else 'same' if same else 'changed'
        counts[status] += 1
        counts['missing'] += row['value'] is None
        if row['value'] is not None and row['variable'] in ('precipitation', 'flow') and row['value'] < 0:
            raise ValueError(f'{parsed[index][1]} {row["sheet"]}/{row["cell"]} 出现负值，请核对原表')
        if row['value'] is not None and row['variable'] == 'precipitation' and row['value'] >= 300 and len(issues) < 30:
            issues.append({'file': parsed[index][1], 'station': row['station'], 'time': row['time'],
                           'value': row['value'], 'message': '降水值偏大，请核对原表'})
        if status == 'changed' and len(changes) < 30:
            changes.append({'station': key[0], 'variable': key[1], 'time': key[3],
                            'old': prior['value'], 'new': row['value'], 'file': parsed[index][1]})
        groups[(index, row['station'], row['variable'], row['step'])].append(row)
    summary = []
    for (index, station, variable, step), rows in groups.items():
        summary.append({'file': parsed[index][1], 'station': station, 'variable': variable, 'step': step,
                        'start': min(r['time'] for r in rows), 'end': max(r['time'] for r in rows),
                        'records': len(rows)})
    files_identity = [{'name': name, 'sha256': sha_file(path), 'station': station, 'variable': variable}
                      for path, name, station, variable, _ in parsed]
    # Include prior active values so a simultaneous update cannot be silently overwritten.
    previous = [{'key': key, 'value': existing[key]['value'], 'quality': existing[key]['quality'],
                 'source': existing[key]['source_id'], 'id': existing[key]['id']}
                for key in sorted(combined) if key in existing]
    token = digest({'files': files_identity, 'previous': previous})
    return parsed, combined, existing, {'files': files_identity, 'summary': summary, 'counts': counts,
                                         'changes': changes, 'issues': issues, 'previewToken': token}


def preview(db, request):
    with db.connect() as conn:
        return _prepare(db, request, conn)[3]


def commit(db, request):
    with db.connect() as conn:
        conn.execute('BEGIN IMMEDIATE')
        parsed, combined, existing, report = _prepare(db, request, conn)
        if request.get('previewToken') != report['previewToken']:
            raise ValueError('资料或现有记录已变化，请重新检查后更新')
        by_file = defaultdict(list)
        for key, (index, row) in combined.items():
            prior = existing.get(key)
            if prior is not None and prior['value'] is not None and row['quality'] == 'missing':
                continue
            if prior is not None and all(prior[field] == value for field, value in (
                    ('value', row['value']), ('quality', row['quality']), ('unit', UNITS[row['variable']]),
                    ('period_start', row['start']), ('period_end', row['end']))):
                continue
            by_file[index].append(row)
        sources = []
        for index, rows in by_file.items():
            path, original, station, variable, _ = parsed[index]
            checksum = sha_file(path)
            identity = digest({'file': checksum, 'station': station, 'variable': variable,
                               'base': report['previewToken'],
                               'rows': [(r['station'], r['step'], r['time'], r['value'], r['quality']) for r in rows]})
            identifier = 'SRC-' + identity[:24]
            logical = 'native-excel/' + identifier
            archive = db.workspace / 'observation-sources' / identifier / original
            archive.parent.mkdir(parents=True, exist_ok=True)
            if not archive.exists():
                with archive.open('xb') as dst, path.open('rb') as src:
                    shutil.copyfileobj(src, dst)
            if sha_file(archive) != checksum:
                raise ValueError('原始文件归档摘要不匹配')
            metadata = {'adapter': 'native-excel-increment-v1', 'fileSHA256': checksum,
                        'archivedPath': archive.relative_to(db.workspace).as_posix(),
                        'originalName': original, 'timeZone': 'Asia/Shanghai',
                        'previousRecordsPreserved': True}
            conn.execute('INSERT INTO sources VALUES(?,?,?,?,?,?)',
                         (identifier, logical, checksum, path.stat().st_size, now(), canonical(metadata)))
            groups = defaultdict(list)
            for row in rows:
                conn.execute('''INSERT INTO observations(source_id,station_id,variable,time,period_start,
                    period_end,step,unit,value,raw_text,quality,sheet,cell) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                    (identifier, row['station'], row['variable'], row['time'], row['start'], row['end'],
                     row['step'], UNITS[row['variable']], row['value'], row['raw'], row['quality'],
                     row['sheet'], row['cell']))
                groups[(row['station'], row['variable'], row['step'])].append(row)
            for (sid, variable, step), block in groups.items():
                conn.execute('INSERT INTO coverage VALUES(?,?,?,?,?,?,?,?,?)',
                    (identifier, sid, variable, step, min(r['time'] for r in block), max(r['time'] for r in block),
                     sum(r['value'] is not None for r in block), len(block),
                     canonical({'coverageConfirmed': False, 'zeroForAbsent': False,
                                'policy': 'explicit-records-only; missing intervals are not zero'})))
            sources.append(identifier)
        conn.execute('INSERT INTO audit(at,action,payload) VALUES(?,?,?)',
                     (now(), 'observations.native-increment', canonical({'sources': sources, **report['counts']})))
    return {'status': 'imported' if sources else 'unchanged', 'sources': sources, **report['counts']}
